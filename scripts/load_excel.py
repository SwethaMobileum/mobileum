import os
import sys
import json
import re
import unicodedata
import argparse
import psycopg2
from psycopg2.extras import execute_values, Json
import openpyxl

def clean_text(s):
    if s is None:
        return None
    val = str(s).replace('\xa0', ' ').replace('\u200b', ' ')
    val = re.sub(r'\s+', ' ', val).strip()
    return val

def normalize_country_name(c):
    if not c:
        return c
    clean_c = clean_text(c)
    low = clean_c.lower()
    if low in ['united arab emirates', 'united arab emirates - africa', 'uae']:
        return 'UAE'
    if low in ['usa', 'united states of america', 'united states']:
        return 'United States'
    if low in ['uk', 'united kingdom']:
        return 'United Kingdom'
    if low in ['ivory coast', "cote d'ivoire"]:
        return "Cote d'Ivoire"
    if low in ['congo', 'dr congo', 'democratic republic of the congo']:
        return 'DR Congo'
    return clean_c

def load_env_local():
    env_paths = ['.env.local', 'src/.env.local', '.env']
    for path in env_paths:
        if os.path.exists(path):
            with open(path, 'r', encoding='utf-8') as f:
                for line in f:
                    line = line.strip()
                    if line and not line.startswith('#') and '=' in line:
                        k, v = line.split('=', 1)
                        if k.strip() in ['DATABASE_URL', 'POSTGRES_URL']:
                            os.environ[k.strip()] = v.strip().strip('"').strip("'")

# Known aliases connecting the 37 canonical health customers with website master telecom operator names
KNOWN_WEBSITE_OPERATOR_MAPPINGS = [
    ("Verizon USA", "Verizon"),
    ("Verizon USA", "Verizon Wireless"),
    ("T-Mobile USA", "T-Mobile US"),
    ("T-Mobile USA", "T-Mobile"),
    ("AT&T USA", "AT&T Mobility"),
    ("AT&T USA", "AT&T"),
    ("Dish USA", "Dish Wireless"),
    ("Dish USA", "Dish"),
    ("Syniverse USA", "Syniverse"),
    ("STC Saudi Arabia", "STC"),
    ("Mobily Saudi Arabia", "Mobily"),
    ("NTT Docomo Japan", "NTT Docomo"),
    ("SoftBank Japan", "SoftBank"),
    ("Rakuten Japan", "Rakuten Mobile"),
    ("Unitel Angola", "Unitel Angola"),
    ("Unitel Angola", "Unitel"),
    ("Deutsche Telekom Germany", "Deutsche Telekom_T-Mobile group"),
    ("Deutsche Telekom Germany", "Telekom Deutschland"),
    ("Deutsche Telekom Germany", "Deutsche Telekom"),
    ("Telefonica Germany", "Telefonica Deutschland GmbH"),
    ("Telefonica Germany", "Telefónica Deutschland (O2)"),
    ("Telefonica Germany", "Telefonica Germany"),
    ("NOS Portugal", "NOS Portugal"),
    ("NOS Portugal", "NOS Comunicações"),
    ("Optus Australia", "Optus (Singtel)"),
    ("Optus Australia", "Optus"),
    ("TPG Australia", "TPG Telecom (Vodafone AU)"),
    ("TPG Australia", "TPG Telecom"),
    ("Telstra Australia", "Telstra"),
    ("Telkomsel Indonesia", "Telkomsel"),
    ("Orange France", "Orange France"),
    ("Orange France", "Orange"),
    ("OINIS France", "OINIS France"),
    ("Etisalat UAE", "e& (Etisalat)"),
    ("Etisalat UAE", "Etisalat"),
    ("Etisalat UAE", "e&"),
    ("Du UAE", "du"),
    ("Du UAE", "Du"),
    ("Grameenphone Bangladesh", "Grameenphone (Telenor)"),
    ("Grameenphone Bangladesh", "Grameenphone"),
    ("TELUS Canada", "Telus"),
    ("TELUS Canada", "TELUS Mobility"),
    ("Bell Canada", "Bell Mobility"),
    ("Bell Canada", "Bell"),
    ("True Move Thailand", "True Move H (True Corp)"),
    ("True Move Thailand", "TrueMove H"),
    ("Orange Spain", "Orange Spain"),
    ("Telefonica Spain", "Telefonica Spain"),
    ("Telefonica Spain", "Telefónica España (Movistar)"),
    ("SingTel Singapore", "Singtel"),
    ("SingTel Singapore", "Singtel Mobile"),
    ("Orange Middle East and Africa", "Sonatel (Orange)"),
    ("Orange Middle East and Africa", "Orange MEA"),
    ("Orange Middle East and Africa", "Orange Senegal"),
    ("Taiwan Mobile Taiwan", "Taiwan Mobile"),
    ("TIM Italy", "Telecom Italia Group"),
    ("TIM Italy", "TIM (Telecom Italia)"),
    ("TIM Italy", "TIM"),
    ("Claro Argentina", "Claro Argentina"),
    ("Claro Argentina", "Claro"),
    ("TeliaSonera Sweden", "Telia Company (Group) - Sweden"),
    ("TeliaSonera Sweden", "Telia Sverige"),
    ("TeliaSonera Sweden", "Telia"),
    ("America Movil Group Mexico", "Telcel"),
    ("America Movil Group Mexico", "América Móvil"),
    ("Telcel Mexico", "Telcel"),
    ("Bayobab (MTN GlobalConnect)", "Bayobab (MTN GlobalConnect)"),
    ("Bayobab (MTN GlobalConnect)", "MTN GlobalConnect"),
    ("Bayobab (MTN GlobalConnect)", "MTN")
]

def parse_health():
    file_path = os.path.join('data-source', 'health.xlsx')
    if not os.path.exists(file_path):
        print(f"Warning: {file_path} not found.")
        return [], [], []
    
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb['Top25 Customers Health']
    
    # Find header row
    header_row_idx = None
    headers = []
    for r_idx, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if row and any(isinstance(cell, str) and "Customer Name" in cell for cell in row):
            header_row_idx = r_idx
            # Clean headers: remove leading/trailing spaces and non-breaking spaces
            headers = [clean_text(c) if c else "" for c in row]
            break
            
    if not header_row_idx:
        print("Header 'Customer Name' not found in health sheet.")
        return [], [], []

    customers = []
    business_units = []
    
    for row in sheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if not row:
            continue
        row_dict = dict(zip(headers, row))
        
        customer_name = clean_text(row_dict.get("Customer Name"))
        if not customer_name:
            continue
            
        rank = row_dict.get("Rank")
        
        # Raw Data parsing
        def safe_val(val):
            if val == "#N/A" or val is None:
                return None
            if isinstance(val, str):
                cleaned = clean_text(val)
                if cleaned == "#N/A" or cleaned == "":
                    return None
                return cleaned
            return val
            
        is_customer = rank is not None and str(rank).strip() != ""
        
        # Determine CSAT color
        csat = safe_val(row_dict.get("CSAT color"))
        if csat and isinstance(csat, str):
            c_low = csat.lower().strip()
            if "bl" in c_low and "e" in c_low: csat = "Blue"
            elif "gr" in c_low: csat = "Green"
            elif "am" in c_low: csat = "Amber"
            elif "re" in c_low: csat = "Red"
            else: csat = None
            
        health_score = safe_val(row_dict.get("Customer Health"))
        status = None
        if health_score is not None:
            if health_score >= 80: status = 'Green'
            elif health_score >= 50: status = 'Amber'
            else: status = 'Red'
            
        country_raw = safe_val(row_dict.get("Country"))
        country_norm = normalize_country_name(country_raw)

        record = {
            "customer_name": customer_name,
            "is_customer": is_customer,
            "group": safe_val(row_dict.get("Group")),
            "country": country_norm,
            "raw_country": country_raw,
            "bu_name_or_count": safe_val(row_dict.get("Business Units")),
            "revenue": safe_val(row_dict.get("2025 Revenues")),
            "health": health_score,
            "behavioral": safe_val(row_dict.get("Behavioral Customer Health")),
            "financial": safe_val(row_dict.get("Financial Customer Health")),
            "sentiment": safe_val(row_dict.get("Sentiment Customer Health")),
            "csat_color": csat,
            "status": status,
            "raw_data": {},
            "scores": {}
        }
        
        col_counts = {}
        for idx, h in enumerate(headers):
            if not h: continue
            col_counts[h] = col_counts.get(h, 0) + 1
            val = safe_val(row[idx])
            
            if h not in ["Rank", "Customer Name", "Group", "Country", "Business Units", "2025 Revenues", "Customer Health", "Behavioral Customer Health", "Financial Customer Health", "Sentiment Customer Health", "CSAT color"]:
                if col_counts[h] == 1:
                    record["raw_data"][h] = val
                elif col_counts[h] == 2:
                    record["scores"][h] = val

        if is_customer:
            customers.append(record)
        else:
            business_units.append(record)
            
    # KPI Definitions
    kpi_sheet = wb['KPI Weighting & Scoring']
    kpi_defs = []
    for row in kpi_sheet.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            kpi_name = clean_text(row[0])
            if kpi_name:
                kpi_defs.append({
                    "kpi_name": kpi_name,
                    "part": safe_val(row[1]),
                    "weight": safe_val(row[2]),
                    "rule": safe_val(row[3])
                })
            
    return customers, business_units, kpi_defs

def parse_mapping():
    file_path = os.path.join('data-source', 'mapping.xlsx')
    if not os.path.exists(file_path):
        return [], []
    wb = openpyxl.load_workbook(file_path, data_only=True)
    
    cust_map = []
    existing_pairs = set()

    if 'Customer_Identifiers' in wb.sheetnames:
        sheet = wb['Customer_Identifiers']
        headers = [clean_text(c) if c else "" for c in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]: continue
            d = dict(zip(headers, row))
            lc = clean_text(d.get("LogicalCustomer", ""))
            val = clean_text(d.get("Value", ""))
            if not lc or not val: continue
            
            existing_pairs.add((lc.lower(), val.lower()))
            cust_map.append({
                "logical_customer": lc,
                "file_name": clean_text(d.get("FileName")),
                "sheet": clean_text(d.get("Sheet")),
                "field": clean_text(d.get("Field")),
                "value": val,
                "customer_code": clean_text(d.get("Customer Code"))
            })
                
    # Inject known website operator mappings if not already present
    for lc, val in KNOWN_WEBSITE_OPERATOR_MAPPINGS:
        pair = (lc.lower(), val.lower())
        if pair not in existing_pairs:
            existing_pairs.add(pair)
            cust_map.append({
                "logical_customer": lc,
                "file_name": "Website Master Telecom",
                "sheet": "Operators",
                "field": "Operator Name",
                "value": val,
                "customer_code": None
            })
            
    bu_map = []
    if 'Business_Unit_Identifiers' in wb.sheetnames:
        sheet = wb['Business_Unit_Identifiers']
        headers = [clean_text(c) if c else "" for c in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]: continue
            d = dict(zip(headers, row))
            lbu = clean_text(d.get("LogicalBusinessUnit", ""))
            val = clean_text(d.get("Value", ""))
            if not lbu or not val: continue
            
            bu_map.append({
                "logical_bu": lbu,
                "file_name": clean_text(d.get("FileName")),
                "sheet": clean_text(d.get("Sheet")),
                "field": clean_text(d.get("Field")),
                "value": val
            })
            
    return cust_map, bu_map

def parse_skygeni():
    file_path = os.path.join('data-source', 'skygeni.xlsx')
    if not os.path.exists(file_path):
        return [], "Unknown Date"
        
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb['Sheet1']
    
    updated_on = "Unknown Date"
    for r_idx in range(1, 7):
        val = sheet.cell(row=r_idx, column=1).value
        if val and isinstance(val, str) and "Updated on:" in val:
            updated_on = clean_text(val)
            break
            
    skygeni_data = []
    header_row_idx = 2
    for r_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=True), 1):
        row_strs = [clean_text(c).lower() if c else "" for c in row]
        if "region" in row_strs and "country" in row_strs:
            header_row_idx = r_idx
            headers = [clean_text(c) if c else "" for c in row]
            break
    else:
        headers = [clean_text(c) if c else "" for c in next(sheet.iter_rows(min_row=7, max_row=7, values_only=True))]
        header_row_idx = 7
        
    for row in sheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if not row or not row[0]: continue
        d = dict(zip(headers, row))
        country_raw = clean_text(d.get("Country"))
        country_norm = normalize_country_name(country_raw)

        raw_owner = clean_text(d.get("Account Owner", d.get("Account_Owner")))
        owner = "Not assigned" if (not raw_owner or raw_owner == "Unknown") else raw_owner

        skygeni_data.append({
            "id": d.get("#"),
            "region": clean_text(d.get("Region")),
            "country": country_norm,
            "raw_country": country_raw,
            "account_name": clean_text(d.get("Account Name", d.get("Account_Name", d.get("Account", "")))),
            "bu": clean_text(d.get("BU", d.get("Business Unit", ""))),
            "product_group": clean_text(d.get("Product Group", d.get("Product_Group"))),
            "account_owner": owner
        })
        
    return skygeni_data, updated_on

def save_json(data, filename):
    os.makedirs('data-cache', exist_ok=True)
    with open(os.path.join('data-cache', filename), 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=2)

def load_to_postgres(customers, business_units, cust_map, bu_map, skygeni_data, kpi_defs, updated_on):
    load_env_local()
    db_url = os.environ.get('DATABASE_URL')
    if not db_url:
        print("DATABASE_URL not found.")
        sys.exit(1)
        
    conn = psycopg2.connect(db_url)
    cur = conn.cursor()
    
    cur.execute("""
    CREATE TABLE IF NOT EXISTS customer_health (
        customer_name TEXT PRIMARY KEY,
        customer_group TEXT,
        country TEXT,
        bu_count NUMERIC,
        revenue TEXT,
        health NUMERIC,
        behavioral NUMERIC,
        financial NUMERIC,
        sentiment NUMERIC,
        csat_color TEXT,
        status TEXT,
        raw_data JSONB,
        scores JSONB
    );
    CREATE TABLE IF NOT EXISTS customer_health_bu (
        id SERIAL PRIMARY KEY,
        customer_name TEXT,
        bu_name TEXT,
        revenue TEXT,
        health NUMERIC,
        behavioral NUMERIC,
        financial NUMERIC,
        sentiment NUMERIC,
        csat_color TEXT,
        status TEXT,
        raw_data JSONB,
        scores JSONB
    );
    CREATE TABLE IF NOT EXISTS customer_mapping (
        logical_customer TEXT,
        file_name TEXT,
        sheet TEXT,
        field TEXT,
        value TEXT,
        customer_code TEXT
    );
    CREATE TABLE IF NOT EXISTS business_unit_mapping (
        logical_bu TEXT,
        file_name TEXT,
        sheet TEXT,
        field TEXT,
        value TEXT
    );
    CREATE TABLE IF NOT EXISTS skygeni_product_groups (
        id SERIAL PRIMARY KEY,
        region TEXT,
        country TEXT,
        account_name TEXT,
        bu TEXT,
        product_group TEXT,
        account_owner TEXT
    );
    CREATE TABLE IF NOT EXISTS kpi_definitions (
        kpi_name TEXT,
        part TEXT,
        weight NUMERIC,
        rule TEXT
    );
    CREATE TABLE IF NOT EXISTS data_status (
        key TEXT PRIMARY KEY,
        value TEXT,
        updated_at TIMESTAMPTZ DEFAULT NOW()
    );
    """)
    
    cur.execute("TRUNCATE TABLE customer_health, customer_health_bu, customer_mapping, business_unit_mapping, skygeni_product_groups, kpi_definitions, data_status CASCADE;")
    
    # customer_health
    ch_rows = [(c["customer_name"], c["group"], c["country"], c["bu_name_or_count"], c["revenue"], c["health"], c["behavioral"], c["financial"], c["sentiment"], c["csat_color"], c["status"], Json(c["raw_data"]), Json(c["scores"])) for c in customers]
    execute_values(cur, "INSERT INTO customer_health VALUES %s", ch_rows)
    
    # customer_health_bu
    chb_rows = [(c["customer_name"], c["bu_name_or_count"], c["revenue"], c["health"], c["behavioral"], c["financial"], c["sentiment"], c["csat_color"], c["status"], Json(c["raw_data"]), Json(c["scores"])) for c in business_units]
    execute_values(cur, "INSERT INTO customer_health_bu (customer_name, bu_name, revenue, health, behavioral, financial, sentiment, csat_color, status, raw_data, scores) VALUES %s", chb_rows)
    
    # mapping
    cm_rows = [(c["logical_customer"], c["file_name"], c["sheet"], c["field"], c["value"], c["customer_code"]) for c in cust_map]
    execute_values(cur, "INSERT INTO customer_mapping VALUES %s", cm_rows)
    
    bum_rows = [(c["logical_bu"], c["file_name"], c["sheet"], c["field"], c["value"]) for c in bu_map]
    execute_values(cur, "INSERT INTO business_unit_mapping VALUES %s", bum_rows)
    
    # skygeni
    sk_rows = [(c["region"], c["country"], c["account_name"], c["bu"], c["product_group"], c["account_owner"]) for c in skygeni_data]
    execute_values(cur, "INSERT INTO skygeni_product_groups (region, country, account_name, bu, product_group, account_owner) VALUES %s", sk_rows)
    
    # kpi
    kpi_rows = [(c["kpi_name"], c["part"], c["weight"], c["rule"]) for c in kpi_defs]
    execute_values(cur, "INSERT INTO kpi_definitions VALUES %s", kpi_rows)
    
    # status
    execute_values(cur, "INSERT INTO data_status (key, value) VALUES %s", [("skygeni_updated_on", updated_on)])
    
    conn.commit()
    cur.close()
    conn.close()
    print("Postgres load complete.")

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--to', choices=['json', 'postgres'], required=True)
    args = parser.parse_args()
    
    print("Parsing Excel files with whitespace cleaning and alias enrichment...")
    customers, business_units, kpi_defs = parse_health()
    cust_map, bu_map = parse_mapping()
    skygeni_data, updated_on = parse_skygeni()
    
    if args.to == 'json':
        save_json(customers, 'customers.json')
        save_json(business_units, 'business_units.json')
        save_json(cust_map, 'customer_mapping.json')
        save_json(bu_map, 'bu_mapping.json')
        save_json(skygeni_data, 'skygeni.json')
        save_json(kpi_defs, 'kpi_definitions.json')
        save_json({"updated_on": updated_on}, 'data_status.json')
        print("JSON load complete. All cache files refreshed.")
    elif args.to == 'postgres':
        load_to_postgres(customers, business_units, cust_map, bu_map, skygeni_data, kpi_defs, updated_on)

if __name__ == '__main__':
    main()
